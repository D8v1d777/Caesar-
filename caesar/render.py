import os
import uuid
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

def render_slide(spec: dict, grounded_data: dict = None) -> str:
    """
    Render a PowerPoint slide based on the specification.

    Args:
        spec: Dictionary from generate_slide_spec containing:
            - title: string
            - chart_type: string (one of: "title", "text", "table", "bar", "line", "pie")
            - data: dict or list appropriate for chart_type
            - source_label: string indicating data source
            - notes: optional string for speaker notes
        grounded_data: Optional dictionary from get_grounded_data (not used in rendering
                       but could be used for validation or additional info)

    Returns:
        Path to the generated .pptx file.
    """
    # Create a presentation object
    prs = Presentation()

    # Get the title and other fields
    title_text = spec.get("title", "Untitled Slide")
    chart_type = spec.get("chart_type", "text")
    data = spec.get("data", {})
    source_label = spec.get("source_label", "Data Source")
    notes_text = spec.get("notes", "")

    # Choose layout based on chart_type
    if chart_type == "title":
        # Title slide layout
        slide_layout = prs.slide_layouts[0]  # Title slide
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]

        title.text = title_text
        subtitle.text = source_label  # Use subtitle for source
    else:
        # For other types, use title and content layout
        slide_layout = prs.slide_layouts[1]  # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        content = slide.placeholders[1]

        title.text = title_text

        if chart_type == "text":
            # Text content
            text_frame = content.text_frame
            text_frame.text = data.get("content", "")
        elif chart_type == "table":
            # Table data: expect {"headers": [...], "rows": [[...], ...]}
            headers = data.get("headers", [])
            rows = data.get("rows", [])
            if not headers and not rows:
                # Fallback: show a message
                text_frame = content.text_frame
                text_frame.text = "No data available"
            else:
                # Create a table
                rows_count = len(rows) + 1  # +1 for header row
                cols_count = len(headers) if headers else len(rows[0]) if rows else 0
                if cols_count == 0:
                    table = None
                else:
                    left = Inches(0.5)
                    top = Inches(1.5)
                    width = Inches(9)
                    height = Inches(0.5 + 0.3 * rows_count)  # approximate height
                    table = shape = slide.shapes.add_table(rows_count, cols_count, left, top, width, height).table

                    # Set column widths (optional)
                    for i, width in enumerate([Inches(2)] * cols_count):  # equal width for simplicity
                        table.columns[i].width = width

                    # Add header row
                    for col_idx, header in enumerate(headers):
                        cell = table.cell(0, col_idx)
                        cell.text = str(header)
                        # Make header bold
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                run.font.bold = True

                    # Add data rows
                    for row_idx, row in enumerate(rows):
                        for col_idx, cell_value in enumerate(row):
                            cell = table.cell(row_idx + 1, col_idx)
                            cell.text = str(cell_value)
        elif chart_type in ("bar", "line", "pie"):
            # Chart data: expect {"labels": [...], "datasets": [{"label": "...", "data": [...]}]}
            # For pie, we might only have one dataset and use labels and data directly.
            labels = data.get("labels", [])
            datasets = data.get("datasets", [])

            if not labels:
                # Fallback to text
                text_frame = content.text_frame
                text_frame.text = "No chart data available"
            else:
                # Determine chart type
                if chart_type == "bar":
                    chart_type_enum = XL_CHART_TYPE.COLUMN_CLUSTERED
                elif chart_type == "line":
                    chart_type_enum = XL_CHART_TYPE.LINE
                else:  # pie
                    chart_type_enum = XL_CHART_TYPE.PIE

                # Create chart data
                chart_data = ChartData()
                chart_data.categories = labels

                # Add each dataset as a series
                for dataset in datasets:
                    label = dataset.get("label", "Series")
                    data_vals = dataset.get("data", [])
                    chart_data.add_series(label, data_vals)

                # If only one dataset and no label provided, we might just use the first series
                # but the above loop works.

                # Add the chart
                left = Inches(0.5)
                top = Inches(1.5)
                width = Inches(9)
                height = Inches(5)
                graphic_frame = slide.shapes.add_chart(
                    chart_type_enum, left, top, width, height, chart_data
                )
                chart = graphic_frame.chart

                # Optional: configure chart properties
                chart.has_legend = True
                chart.legend.position = XL_LEGEND_POSITION.TOP
                chart.has_title = True
                chart.chart_title.text_frame.text = title_text
        else:
            # Fallback to text
            text_frame = content.text_frame
            text_frame.text = f"Unsupported chart type: {chart_type}"

    # Add footer with source label (optional)
    # We'll add a textbox at the bottom
    left = Inches(0.5)
    top = Inches(6.5)  # adjust based on slide size
    width = Inches(9)
    height = Inches(0.5)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.text = f"Source: {source_label}"
    # Format the footer text
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(10)
            run.font.italic = True

    # Add notes to the slide
    if notes_text:
        notes_slide = slide.notes_slide
        notes_text_frame = notes_slide.notes_text_frame
        notes_text_frame.text = notes_text

    # Save the presentation
    output_dir = "generated_slides"
    os.makedirs(output_dir, exist_ok=True)
    filename = f"{uuid.uuid4()}.pptx"
    output_path = os.path.join(output_dir, filename)
    prs.save(output_path)

    return output_path